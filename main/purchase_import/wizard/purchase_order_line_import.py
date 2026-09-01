import base64
import binascii
import hashlib
import io
import json
import mimetypes
import os

from odoo import Command, _, api, fields, models
from odoo.exceptions import AccessError, UserError, ValidationError
from odoo.addons.base_import.models.base_import import ImportValidationError

from odoo.addons.purchase_import.models.import_profile import (
    LOOKUP_SELECTION,
    TARGET_SELECTION,
)


class PurchaseOrderLineImportWizard(models.TransientModel):
    _name = "purchase.order.line.import.wizard"
    _description = "Purchase Order Line Import"
    _check_company_auto = True

    order_id = fields.Many2one(
        "purchase.order", required=True, readonly=True, ondelete="cascade"
    )
    company_id = fields.Many2one(related="order_id.company_id")
    profile_id = fields.Many2one(
        "purchase.import.profile",
        check_company=True,
        domain="[('company_id', '=', company_id), ('active', '=', True)]",
    )
    profile_name = fields.Char(help="Name used when saving the current settings as a new profile.")
    lookup_order = fields.Char(
        default="supplier_code,default_code,barcode",
        readonly=True,
    )
    file_data = fields.Binary(required=True, attachment=False)
    file_name = fields.Char(required=True)
    encoding = fields.Char()
    separator = fields.Char()
    quoting = fields.Char(default='"', required=True)
    decimal_separator = fields.Selection(
        [("dot", "Dot"), ("comma", "Comma")], default="dot", required=True
    )
    thousand_separator = fields.Selection(
        [("none", "None"), ("comma", "Comma"), ("dot", "Dot"), ("space", "Space")],
        default="none",
        required=True,
    )
    sheet = fields.Char()
    available_sheets = fields.Char(readonly=True)
    header_row = fields.Integer(default=1, required=True)
    data_row = fields.Integer(default=2, required=True)
    max_rows = fields.Integer(default=5000, required=True)
    max_file_size_mb = fields.Integer(default=10, required=True)
    unmatched_policy = fields.Selection(
        [("reject", "Reject"), ("create", "Create Product")],
        default="reject",
        required=True,
    )
    creation_mode = fields.Selection(
        [("standalone", "Standalone Product"), ("variant", "Template Variant")],
        default="standalone",
        required=True,
    )
    variant_template_id = fields.Many2one(
        "product.template",
        check_company=True,
        domain="[('company_id', 'in', [False, company_id])]",
    )
    default_product_type = fields.Selection(
        [("consu", "Goods"), ("service", "Service")],
        default="consu",
        required=True,
    )
    default_is_storable = fields.Boolean()
    default_categ_id = fields.Many2one("product.category")
    default_uom_id = fields.Many2one("uom.uom")
    default_purchase_uom_id = fields.Many2one("uom.uom")
    default_supplier_tax_ids = fields.Many2many(
        "account.tax",
        "purchase_import_wizard_supplier_tax_rel",
        "wizard_id",
        "tax_id",
        domain="[('type_tax_use', '=', 'purchase'), ('company_id', '=', company_id)]",
        check_company=True,
    )
    default_tracking = fields.Selection(
        [("none", "By Quantity"), ("lot", "By Lots"), ("serial", "By Serial Number")],
        default="none",
        required=True,
    )
    default_purchase_ok = fields.Boolean(default=True)
    default_sale_ok = fields.Boolean(default=False)
    mapping_ids = fields.One2many(
        "purchase.order.line.import.mapping", "wizard_id"
    )
    preview_ids = fields.One2many(
        "purchase.order.line.import.preview", "wizard_id", readonly=True
    )
    state = fields.Selection(
        [("upload", "Upload"), ("mapping", "Mapping"), ("preview", "Preview")],
        default="upload",
        required=True,
        readonly=True,
    )
    preview_valid = fields.Boolean(readonly=True)
    source_row_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)
    validation_token = fields.Char(readonly=True)
    order_snapshot = fields.Char(readonly=True)

    @api.onchange("profile_id")
    def _onchange_profile_id(self):
        if not self.profile_id:
            return
        values = self._profile_values(self.profile_id)
        for field_name, value in values.items():
            self[field_name] = value
        self.lookup_order = ",".join(
            self.profile_id.lookup_ids.sorted("sequence").mapped("lookup_type")
        )
        self.mapping_ids = [Command.clear()] + [
            Command.create({
                "sequence": mapping.sequence,
                "source_index": index,
                "source_column": mapping.source_column,
                "target": mapping.target,
                "attribute_id": mapping.attribute_id.id,
            })
            for index, mapping in enumerate(self.profile_id.mapping_ids)
        ]
        self.preview_valid = False
        self.state = "upload"

    @api.constrains("header_row", "data_row", "max_rows", "max_file_size_mb")
    def _check_rows_and_limits(self):
        for wizard in self:
            if wizard.header_row <= 0 or wizard.data_row <= wizard.header_row:
                raise ValidationError(_("The first data row must be after the header row."))
            if wizard.max_rows <= 0 or wizard.max_file_size_mb <= 0:
                raise ValidationError(_("File size and row limits must be positive."))

    @api.constrains("separator", "quoting", "decimal_separator", "thousand_separator")
    def _check_format_characters(self):
        for wizard in self:
            if wizard.separator and len(wizard.separator) != 1:
                raise ValidationError(_("The CSV separator must be one character."))
            if len(wizard.quoting) != 1:
                raise ValidationError(_("The text delimiter must be one character."))
            if wizard.thousand_separator != "none" and (
                wizard.thousand_separator == wizard.decimal_separator
            ):
                raise ValidationError(_("Decimal and thousand separators must be different."))

    @api.constrains(
        "default_product_type", "default_is_storable", "default_tracking"
    )
    def _check_product_defaults(self):
        for wizard in self:
            if wizard.default_product_type == "service" and (
                wizard.default_is_storable or wizard.default_tracking != "none"
            ):
                raise ValidationError(_("A service product cannot be storable or tracked."))
            if not wizard.default_is_storable and wizard.default_tracking != "none":
                raise ValidationError(_("Tracking requires a storable product."))

    def write(self, vals):
        watched_fields = {
            "profile_id",
            "lookup_order",
            "file_data",
            "file_name",
            "encoding",
            "separator",
            "quoting",
            "decimal_separator",
            "thousand_separator",
            "sheet",
            "header_row",
            "data_row",
            "max_rows",
            "max_file_size_mb",
            "unmatched_policy",
            "creation_mode",
            "variant_template_id",
            "default_product_type",
            "default_is_storable",
            "default_categ_id",
            "default_uom_id",
            "default_purchase_uom_id",
            "default_supplier_tax_ids",
            "default_tracking",
            "default_purchase_ok",
            "default_sale_ok",
        }
        if (
            not self.env.context.get("purchase_import_keep_preview")
            and watched_fields & vals.keys()
        ):
            vals = {
                **vals,
                "preview_valid": False,
                "validation_token": False,
                "state": "mapping" if any(wizard.mapping_ids for wizard in self) else "upload",
            }
        return super().write(vals)

    def _profile_values(self, profile):
        return {
            field_name: profile[field_name].id
            if profile._fields[field_name].type == "many2one"
            else profile[field_name].ids
            if profile._fields[field_name].type == "many2many"
            else profile[field_name]
            for field_name in (
                "encoding",
                "separator",
                "quoting",
                "decimal_separator",
                "thousand_separator",
                "sheet",
                "header_row",
                "data_row",
                "max_rows",
                "max_file_size_mb",
                "unmatched_policy",
                "creation_mode",
                "variant_template_id",
                "default_product_type",
                "default_is_storable",
                "default_categ_id",
                "default_uom_id",
                "default_purchase_uom_id",
                "default_supplier_tax_ids",
                "default_tracking",
                "default_purchase_ok",
                "default_sale_ok",
            )
        }

    def _check_order(self):
        self.ensure_one()
        self.order_id._check_line_import_allowed()
        if self.profile_id and self.profile_id.company_id != self.order_id.company_id:
            raise UserError(_("The import profile must belong to the purchase order company."))

    def _get_raw_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Upload a CSV or XLSX file."))
        value = bytes(self.file_data)
        try:
            raw = base64.b64decode(value, validate=True)
        except (binascii.Error, ValueError):
            raw = value
        if len(raw) > self.max_file_size_mb * 1024 * 1024:
            raise UserError(_("The import file exceeds the configured size limit."))
        return raw

    def _parse_file(self):
        self.ensure_one()
        raw = self._get_raw_file()
        extension = os.path.splitext(self.file_name or "")[1].casefold()
        if extension not in {".csv", ".xlsx"}:
            raise UserError(_("Only CSV and XLSX files are supported."))
        file_type = mimetypes.guess_type(self.file_name or "")[0] or "application/octet-stream"
        options = {
            "encoding": self.encoding or "",
            "separator": self.separator or "",
            "quoting": self.quoting,
            "sheet": self.sheet or "",
            "sheets": [],
        }
        try:
            if extension == ".xlsx":
                result = self._read_xlsx_file(raw, options)
            else:
                importer = self.env["base_import.import"].create({
                    "res_model": "purchase.order.line",
                    "file": raw,
                    "file_name": self.file_name,
                    "file_type": file_type,
                })
                result = importer._read_file(options)
        except (ImportValidationError, KeyError, ValueError) as error:
            raise UserError(
                _("The import file could not be read: %(message)s", message=str(error))
            ) from error
        if not result:
            raise UserError(_("The import file has no content."))
        file_length, rows = result
        if file_length <= 0 or not rows:
            raise UserError(_("The import file has no content."))
        if self.header_row > len(rows):
            raise UserError(_("The configured header row does not exist in the file."))
        raw_headers = [str(value).strip() for value in rows[self.header_row - 1]]
        source_indices = [
            index for index, header in enumerate(raw_headers) if header
        ]
        headers = [raw_headers[index] for index in source_indices]
        if not headers:
            raise UserError(_("Every source column must have a non-empty header."))
        normalized_headers = [self._normalize_header(header) for header in headers]
        if len(set(normalized_headers)) != len(normalized_headers):
            raise UserError(_("Source column headers must be unique."))
        normalized_rows = []
        row_numbers = []
        for row_number, row in enumerate(
            rows[self.data_row - 1 :], start=self.data_row
        ):
            if not any(self._has_value(value) for value in row):
                break
            if len(row) > len(raw_headers):
                raise UserError(
                    _(
                        "Source row %(row)s has more values than the header.",
                        row=row_number,
                    )
                )
            normalized_row = [
                row[index] if index < len(row) else ""
                for index in source_indices
            ]
            if any(self._has_value(value) for value in normalized_row):
                normalized_rows.append(normalized_row)
                row_numbers.append(row_number)
        if not normalized_rows:
            raise UserError(_("The selected data range contains no import rows."))
        if len(normalized_rows) > self.max_rows:
            raise UserError(_("The import file exceeds the configured row limit."))
        return {
            "headers": headers,
            "rows": normalized_rows,
            "row_numbers": row_numbers,
            "sheet": options.get("sheet") or "",
            "sheets": options.get("sheets") or [],
        }

    @api.model
    def _read_xlsx_file(self, raw, options):
        import openpyxl  # noqa: PLC0415
        import openpyxl.cell.cell as cell_types  # noqa: PLC0415
        import openpyxl.styles.numbers as styles  # noqa: PLC0415

        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        options["sheets"] = book.sheetnames
        sheet_name = options["sheet"] = options.get("sheet") or book.sheetnames[0]
        sheet = book[sheet_name]
        rows = []
        for row_number, row in enumerate(sheet.rows, 1):
            values = []
            for column_number, cell in enumerate(row, 1):
                if cell.data_type == cell_types.TYPE_ERROR:
                    raise ValueError(
                        _(
                            "Invalid cell value at row %(row)s, column %(column)s: "
                            "%(value)s",
                            row=row_number,
                            column=column_number,
                            value=cell.value,
                        )
                    )
                if cell.value is None:
                    values.append("")
                elif isinstance(cell.value, float):
                    values.append(
                        str(int(cell.value))
                        if cell.value % 1 == 0
                        else str(cell.value)
                    )
                elif cell.is_date:
                    date_format = styles.is_datetime(cell.number_format)
                    values.append(
                        cell.value.date() if date_format == "date" else cell.value
                    )
                else:
                    values.append(str(cell.value))
            rows.append(values)
        return sheet.max_row, rows

    @api.model
    def _has_value(self, value):
        if value is None or value is False:
            return False
        return bool(value.strip()) if isinstance(value, str) else True

    @api.model
    def _normalize_header(self, value):
        return " ".join(str(value).strip().casefold().split())

    def action_parse(self):
        self.ensure_one()
        self._check_order()
        parsed = self._parse_file()
        profile_mappings = {
            self._normalize_header(mapping.source_column): mapping
            for mapping in self.profile_id.mapping_ids
        }
        commands = [Command.clear()]
        for index, header in enumerate(parsed["headers"]):
            profile_mapping = profile_mappings.get(self._normalize_header(header))
            commands.append(Command.create({
                "sequence": (index + 1) * 10,
                "source_index": index,
                "source_column": header,
                "target": profile_mapping.target if profile_mapping else False,
                "attribute_id": profile_mapping.attribute_id.id if profile_mapping else False,
            }))
        self.with_context(purchase_import_keep_preview=True).write({
            "mapping_ids": commands,
            "preview_ids": [Command.clear()],
            "available_sheets": ", ".join(parsed["sheets"]),
            "sheet": parsed["sheet"],
            "state": "mapping",
            "preview_valid": False,
            "validation_token": False,
            "source_row_count": len(parsed["rows"]),
            "error_count": 0,
        })
        return self._action_reopen()

    def action_save_profile(self):
        self.ensure_one()
        if not self.env.user.has_group("purchase.group_purchase_manager"):
            raise AccessError(_("Only Purchase Managers can save import profiles."))
        if not self.profile_name:
            raise UserError(_("Enter a name for the new import profile."))
        self._validate_mapping()
        lookup_types = self._lookup_types()
        profile = self.env["purchase.import.profile"].create({
            "name": self.profile_name,
            "company_id": self.order_id.company_id.id,
            **self._wizard_profile_create_values(),
            "mapping_ids": [Command.create({
                "sequence": mapping.sequence,
                "source_column": mapping.source_column,
                "target": mapping.target,
                "attribute_id": mapping.attribute_id.id,
            }) for mapping in self.mapping_ids.filtered("target")],
            "lookup_ids": [Command.create({
                "sequence": sequence,
                "lookup_type": lookup_type,
            }) for sequence, lookup_type in enumerate(lookup_types, start=10)],
        })
        self.profile_id = profile
        return self._action_reopen()

    def _wizard_profile_create_values(self):
        return {
            field_name: self[field_name].id
            if self._fields[field_name].type == "many2one"
            else [Command.set(self[field_name].ids)]
            if self._fields[field_name].type == "many2many"
            else self[field_name]
            for field_name in (
                "encoding",
                "separator",
                "quoting",
                "decimal_separator",
                "thousand_separator",
                "sheet",
                "header_row",
                "data_row",
                "max_rows",
                "max_file_size_mb",
                "unmatched_policy",
                "creation_mode",
                "variant_template_id",
                "default_product_type",
                "default_is_storable",
                "default_categ_id",
                "default_uom_id",
                "default_purchase_uom_id",
                "default_supplier_tax_ids",
                "default_tracking",
                "default_purchase_ok",
                "default_sale_ok",
            )
        }

    def _validate_mapping(self):
        self.ensure_one()
        mappings = self.mapping_ids.filtered("target")
        targets = mappings.filtered(lambda mapping: mapping.target != "attribute").mapped("target")
        duplicates = {target for target in targets if targets.count(target) > 1}
        if duplicates:
            raise UserError(_("Each non-attribute target can only be mapped once."))
        attribute_ids = mappings.filtered(
            lambda mapping: mapping.target == "attribute"
        ).attribute_id.ids
        if len(attribute_ids) != len(set(attribute_ids)):
            raise UserError(_("Each variant attribute can only be mapped once."))
        missing = {"quantity", "price"} - set(targets)
        if missing:
            raise UserError(_("Quantity and unit price columns are required."))
        lookup_targets = {"supplier_code", "default_code", "barcode"} & set(targets)
        has_attributes = bool(attribute_ids)
        if not lookup_targets and not (
            self.unmatched_policy == "create"
            and (
                (self.creation_mode == "variant" and has_attributes)
                or (self.creation_mode == "standalone" and "product_name" in targets)
            )
        ):
            raise UserError(
                _("Map a product identifier, a product name, or a complete variant combination.")
            )
        if (
            self.unmatched_policy == "create"
            and self.creation_mode == "standalone"
            and "product_name" not in targets
        ):
            raise UserError(_("Product name must be mapped for standalone product creation."))
        if (
            self.unmatched_policy == "create"
            and self.creation_mode == "variant"
            and not self.variant_template_id
        ):
            raise UserError(_("Select a product template for variant creation."))

    def _parse_float(self, value, label, required=False):
        value = str(value or "").strip()
        if not value:
            if required:
                raise ValidationError(_("%s is required.", label))
            return 0.0
        separators = {"none": "", "comma": ",", "dot": ".", "space": " "}
        thousand_separator = separators[self.thousand_separator]
        decimal_separator = separators[self.decimal_separator]
        if thousand_separator:
            value = value.replace(thousand_separator, "")
        if decimal_separator != ".":
            value = value.replace(decimal_separator, ".")
        try:
            return float(value)
        except ValueError as error:
            raise ValidationError(_("%s must be a number.", label)) from error

    def _row_values(self, row):
        values = {}
        attributes = {}
        for mapping in self.mapping_ids.filtered("target"):
            value = row[mapping.source_index] if mapping.source_index < len(row) else ""
            if isinstance(value, str):
                value = value.strip()
            if mapping.target == "attribute":
                attributes[mapping.attribute_id.id] = value
            else:
                values[mapping.target] = value
        values["attributes"] = attributes
        return values

    def _lookup_types(self):
        allowed = {item[0] for item in LOOKUP_SELECTION}
        lookup_types = [
            value.strip() for value in (self.lookup_order or "").split(",") if value.strip()
        ]
        if not lookup_types:
            return [item[0] for item in LOOKUP_SELECTION]
        if len(lookup_types) != len(set(lookup_types)) or not set(lookup_types) <= allowed:
            raise ValidationError(_("The product lookup order is invalid."))
        return lookup_types

    def _build_lookup_cache(self, row_values):
        Product = self.env["product.product"].with_company(self.order_id.company_id)
        company_domain = [
            ("product_tmpl_id.company_id", "in", [False, self.order_id.company_id.id])
        ]
        cache = {}
        for lookup_type in self._lookup_types():
            identifiers = {
                values.get(lookup_type) for values in row_values if values.get(lookup_type)
            }
            if not identifiers:
                continue
            if lookup_type == "supplier_code":
                supplierinfos = self.env["product.supplierinfo"].search([
                    ("partner_id", "=", self.order_id.partner_id.commercial_partner_id.id),
                    ("product_code", "in", list(identifiers)),
                    ("company_id", "in", [False, self.order_id.company_id.id]),
                ])
                for supplierinfo in supplierinfos:
                    key = (lookup_type, supplierinfo.product_code)
                    products = (
                        supplierinfo.product_id
                        or supplierinfo.product_tmpl_id.product_variant_ids
                    )
                    cache[key] = (cache.get(key, Product) | products).filtered(
                        lambda product: product.active
                        and (
                            not product.company_id
                            or product.company_id == self.order_id.company_id
                        )
                    )
            else:
                products = Product.search(
                    company_domain + [(lookup_type, "in", list(identifiers))]
                )
                for product in products:
                    key = (lookup_type, product[lookup_type])
                    cache[key] = cache.get(key, Product) | product
        return cache

    def _resolve_product(self, values, lookup_cache):
        Product = self.env["product.product"].with_company(self.order_id.company_id)
        for lookup_type in self._lookup_types():
            value = values.get(lookup_type)
            if not value:
                continue
            products = lookup_cache.get((lookup_type, value), Product)
            if len(products) > 1:
                raise ValidationError(
                    _("Product identifier %(value)s is ambiguous.", value=value)
                )
            if products:
                return products
        return Product

    def _variant_plan(self, values):
        template = self.variant_template_id.with_company(self.order_id.company_id)
        if template.company_id and template.company_id != self.order_id.company_id:
            raise ValidationError(_("The variant template belongs to another company."))
        defining_lines = template.valid_product_template_attribute_line_ids.filtered(
            lambda line: line.attribute_id.create_variant != "no_variant"
        )
        mapped_values = values["attributes"]
        if set(defining_lines.attribute_id.ids) != set(mapped_values):
            raise ValidationError(_("The row must provide every variant-defining attribute."))
        combination = self.env["product.template.attribute.value"]
        labels = []
        for line in defining_lines:
            raw_value = str(mapped_values[line.attribute_id.id] or "").strip()
            matches = line.product_template_value_ids.filtered(
                lambda ptav: ptav.ptav_active
                and ptav.product_attribute_value_id.name == raw_value
            )
            if len(matches) != 1:
                raise ValidationError(
                    _(
                        "Attribute %(attribute)s value %(value)s is missing or ambiguous.",
                        attribute=line.attribute_id.display_name,
                        value=raw_value,
                    )
                )
            combination |= matches
            labels.append(
                f"{line.attribute_id.display_name}: "
                f"{matches.product_attribute_value_id.display_name}"
            )
        if not template._is_combination_possible_by_config(combination, ignore_no_variant=True):
            raise ValidationError(_("The variant combination is not allowed by the template."))
        product = template._get_variant_for_combination(combination)
        if product and not product.active:
            raise ValidationError(_("The matching product variant is archived."))
        if not product and not template.has_dynamic_attributes():
            raise ValidationError(
                _("A missing variant can only be created for dynamic attributes.")
            )
        return {
            "product": product,
            "plan_kind": False if product else "variant",
            "plan_key": f"variant:{template.id}:{','.join(map(str, sorted(combination.ids)))}",
            "combination_ids": combination.ids,
            "attribute_labels": ", ".join(labels),
        }

    def _standalone_plan(self, values):
        name = str(values.get("product_name") or "").strip()
        if not name:
            raise ValidationError(_("Product name is required to create a standalone product."))
        identity = next(
            (
                f"{key}:{values[key]}"
                for key in ("supplier_code", "default_code", "barcode")
                if values.get(key)
            ),
            f"name:{name.casefold()}",
        )
        return {
            "product": self.env["product.product"],
            "plan_kind": "standalone",
            "plan_key": f"standalone:{identity}",
            "combination_ids": [],
            "attribute_labels": "",
        }

    def _resolve_uom(self, raw_value, product, plan_kind):
        if raw_value:
            uoms = self.env["uom.uom"].search([("name", "=ilike", raw_value)], limit=2)
            if len(uoms) != 1:
                raise ValidationError(_("Unit of measure %s is missing or ambiguous.", raw_value))
            return uoms
        if product:
            return product.uom_id
        if plan_kind == "variant":
            return self.default_purchase_uom_id or self.variant_template_id.uom_id
        return (
            self.default_purchase_uom_id
            or self.default_uom_id
            or self.env.ref("uom.product_uom_unit")
        )

    def _normalize_row(self, values, source_row, lookup_cache):
        quantity = self._parse_float(values.get("quantity"), _("Quantity"), required=True)
        price = self._parse_float(values.get("price"), _("Unit Price"), required=True)
        min_qty = self._parse_float(values.get("min_qty"), _("Minimum Quantity"))
        if quantity <= 0:
            raise ValidationError(_("Quantity must be greater than zero."))
        if price < 0 or min_qty < 0:
            raise ValidationError(_("Price and minimum quantity cannot be negative."))
        product = self._resolve_product(values, lookup_cache)
        if product:
            plan = {
                "product": product,
                "plan_kind": False,
                "plan_key": f"product:{product.id}",
                "combination_ids": [],
                "attribute_labels": "",
            }
        elif self.unmatched_policy == "reject":
            raise ValidationError(_("No product matches the configured lookup rules."))
        elif self.creation_mode == "variant":
            plan = self._variant_plan(values)
        else:
            plan = self._standalone_plan(values)
        if plan["plan_kind"] and not self.env["product.product"].has_access("create"):
            raise AccessError(_("You do not have permission to create products or variants."))
        uom = self._resolve_uom(values.get("uom"), plan["product"], plan["plan_kind"])
        return {
            **values,
            **plan,
            "source_row": source_row,
            "quantity": quantity,
            "price": price,
            "min_qty": min_qty,
            "uom": uom,
            "errors": [],
        }

    def _supplierinfo_key(self, normalized, product):
        return product.id, normalized["uom"].id, normalized["min_qty"]

    def _find_supplierinfos(self, normalized_products):
        SupplierInfo = self.env["product.supplierinfo"].with_company(
            self.order_id.company_id
        )
        if not normalized_products:
            return {}
        supplierinfos = SupplierInfo.search([
            ("partner_id", "=", self.order_id.partner_id.commercial_partner_id.id),
            ("product_id", "in", [product.id for _normalized, product in normalized_products]),
            ("company_id", "=", self.order_id.company_id.id),
            ("currency_id", "=", self.order_id.currency_id.id),
            (
                "product_uom_id",
                "in",
                [normalized["uom"].id for normalized, _product in normalized_products],
            ),
            (
                "min_qty",
                "in",
                [normalized["min_qty"] for normalized, _product in normalized_products],
            ),
        ])
        by_key = {}
        for supplierinfo in supplierinfos:
            key = (
                supplierinfo.product_id.id,
                supplierinfo.product_uom_id.id,
                supplierinfo.min_qty,
            )
            by_key[key] = by_key.get(key, SupplierInfo) | supplierinfo
        return by_key

    def _validate_rows(self, parsed):
        rows_with_numbers = [
            (row_number, self._row_values(row))
            for row_number, row in zip(parsed["row_numbers"], parsed["rows"])
        ]
        rows_with_numbers = [
            (row_number, values)
            for row_number, values in rows_with_numbers
            if any(
                self._has_value(value)
                for key, value in values.items()
                if key != "attributes"
            )
            or any(self._has_value(value) for value in values["attributes"].values())
        ]
        row_values = [values for _row_number, values in rows_with_numbers]
        lookup_cache = self._build_lookup_cache(row_values)
        normalized_rows = []
        for source_row, values in rows_with_numbers:
            try:
                normalized = self._normalize_row(values, source_row, lookup_cache)
            except (AccessError, ValidationError, UserError) as error:
                normalized = {
                    "source_row": source_row,
                    "errors": [str(error)],
                    "product": self.env["product.product"],
                    "plan_kind": False,
                    "plan_key": f"error:{source_row}",
                    "combination_ids": [],
                    "attribute_labels": "",
                    "quantity": 0.0,
                    "price": 0.0,
                    "min_qty": 0.0,
                    "uom": self.env["uom.uom"],
                }
            normalized_rows.append(normalized)
        price_keys = {}
        for normalized in normalized_rows:
            if normalized["errors"]:
                continue
            logical_key = (
                normalized["plan_key"],
                normalized["uom"].id,
                normalized["min_qty"],
            )
            previous = price_keys.get(logical_key)
            if previous and previous["price"] != normalized["price"]:
                message = _("Rows with the same supplier price key contain conflicting prices.")
                previous["errors"].append(message)
                normalized["errors"].append(message)
            else:
                price_keys[logical_key] = normalized
        resolved_rows = [
            (normalized, normalized["product"])
            for normalized in normalized_rows
            if not normalized["errors"] and normalized["product"]
        ]
        supplierinfos_by_key = self._find_supplierinfos(resolved_rows)
        for normalized, product in resolved_rows:
            supplierinfos = supplierinfos_by_key.get(
                self._supplierinfo_key(normalized, product),
                self.env["product.supplierinfo"],
            )
            if len(supplierinfos) > 1:
                normalized["errors"].append(
                    _("Multiple supplier prices already exist for the exact import key.")
                )
        return normalized_rows

    def _preview_values(self, normalized):
        return {
            "source_row": normalized["source_row"],
            "status": "error"
            if normalized["errors"]
            else "create"
            if normalized["plan_kind"]
            else "resolved",
            "product_id": normalized["product"].id,
            "planned_action": normalized["plan_kind"] or "existing",
            "product_name": normalized.get("product_name") or normalized["product"].display_name,
            "quantity": normalized["quantity"],
            "uom_id": normalized["uom"].id,
            "price": normalized["price"],
            "min_qty": normalized["min_qty"],
            "supplier_code": normalized.get("supplier_code"),
            "default_code": normalized.get("default_code"),
            "barcode": normalized.get("barcode"),
            "attribute_values": normalized["attribute_labels"],
            "error_message": "\n".join(normalized["errors"]),
        }

    def action_preview(self):
        self.ensure_one()
        self._check_order()
        self._validate_mapping()
        parsed = self._parse_file()
        normalized_rows = self._validate_rows(parsed)
        error_count = sum(bool(row["errors"]) for row in normalized_rows)
        self.with_context(purchase_import_keep_preview=True).write({
            "preview_ids": [Command.clear()] + [
                Command.create(self._preview_values(row)) for row in normalized_rows
            ],
            "state": "preview",
            "preview_valid": not error_count,
            "source_row_count": len(normalized_rows),
            "error_count": error_count,
            "validation_token": self._configuration_digest(),
            "order_snapshot": self._order_line_snapshot(),
        })
        return self._action_reopen()

    def _configuration_digest(self):
        raw = self._get_raw_file()
        payload = {
            "file": hashlib.sha256(raw).hexdigest(),
            "profile": [self.profile_id.id, fields.Datetime.to_string(self.profile_id.write_date)],
            "lookup_order": self.lookup_order,
            "settings": {
                field_name: self[field_name].id
                if self._fields[field_name].type == "many2one"
                else self[field_name].ids
                if self._fields[field_name].type == "many2many"
                else self[field_name]
                for field_name in (
                    "encoding",
                    "separator",
                    "quoting",
                    "decimal_separator",
                    "thousand_separator",
                    "sheet",
                    "header_row",
                    "data_row",
                    "max_rows",
                    "max_file_size_mb",
                    "unmatched_policy",
                    "creation_mode",
                    "variant_template_id",
                    "default_product_type",
                    "default_is_storable",
                    "default_categ_id",
                    "default_uom_id",
                    "default_purchase_uom_id",
                    "default_supplier_tax_ids",
                    "default_tracking",
                    "default_purchase_ok",
                    "default_sale_ok",
                )
            },
            "mapping": [
                [
                    mapping.source_index,
                    mapping.source_column,
                    mapping.target,
                    mapping.attribute_id.id,
                ]
                for mapping in self.mapping_ids.sorted("sequence")
            ],
        }
        return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode()).hexdigest()

    def _order_line_snapshot(self):
        payload = {
            "order": [
                self.order_id.state,
                self.order_id.locked,
                self.order_id.partner_id.id,
                self.order_id.company_id.id,
                self.order_id.currency_id.id,
            ],
            "lines": [
                [line.id, fields.Datetime.to_string(line.write_date)]
                for line in self.order_id.order_line.sorted("id")
            ],
        }
        return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()

    def _create_standalone_product(self, normalized):
        ProductTemplate = self.env["product.template"].with_company(self.order_id.company_id)
        ProductTemplate.check_access("create")
        self.env["product.product"].check_access("create")
        values = {
            "name": normalized["product_name"],
            "type": self.default_product_type,
            "is_storable": (
                self.default_is_storable
                if self.default_product_type == "consu"
                else False
            ),
            "purchase_ok": self.default_purchase_ok,
            "sale_ok": self.default_sale_ok,
            "tracking": self.default_tracking if self.default_is_storable else "none",
            "company_id": self.order_id.company_id.id,
        }
        if self.default_categ_id:
            values["categ_id"] = self.default_categ_id.id
        if self.default_uom_id:
            values["uom_id"] = self.default_uom_id.id
        if self.default_supplier_tax_ids:
            values["supplier_taxes_id"] = [Command.set(self.default_supplier_tax_ids.ids)]
        if normalized.get("default_code"):
            values["default_code"] = normalized["default_code"]
        if normalized.get("barcode"):
            values["barcode"] = normalized["barcode"]
        return ProductTemplate.create(values).product_variant_id

    def _create_dynamic_variant(self, normalized):
        Product = self.env["product.product"]
        Product.check_access("create")
        combination = self.env["product.template.attribute.value"].browse(
            normalized["combination_ids"]
        ).exists()
        product = self.variant_template_id._create_product_variant(combination)
        if not product:
            raise UserError(_("The product variant could not be created."))
        values = {}
        if normalized.get("default_code"):
            values["default_code"] = normalized["default_code"]
        if normalized.get("barcode"):
            values["barcode"] = normalized["barcode"]
        if values:
            product.check_access("write")
            product.write(values)
        return product

    def _supplierinfo_update_values(self, normalized):
        values = {"price": normalized["price"]}
        if normalized.get("supplier_code"):
            values["product_code"] = normalized["supplier_code"]
        if normalized.get("supplier_name"):
            values["product_name"] = normalized["supplier_name"]
        return values

    def _upsert_supplierinfos(self, normalized_products):
        SupplierInfo = self.env["product.supplierinfo"].with_company(self.order_id.company_id)
        supplierinfos_by_key = self._find_supplierinfos(normalized_products)
        create_values = []
        handled_keys = set()
        for normalized, product in normalized_products:
            key = self._supplierinfo_key(normalized, product)
            if key in handled_keys:
                continue
            handled_keys.add(key)
            supplierinfos = supplierinfos_by_key.get(key, SupplierInfo)
            if len(supplierinfos) > 1:
                raise UserError(_("Multiple supplier prices exist for the exact import key."))
            values = self._supplierinfo_update_values(normalized)
            if supplierinfos:
                supplierinfos.check_access("write")
                supplierinfos.write(values)
                continue
            create_values.append({
                **values,
                "partner_id": self.order_id.partner_id.commercial_partner_id.id,
                "product_id": product.id,
                "product_tmpl_id": product.product_tmpl_id.id,
                "company_id": self.order_id.company_id.id,
                "currency_id": self.order_id.currency_id.id,
                "product_uom_id": normalized["uom"].id,
                "min_qty": normalized["min_qty"],
            })
        if create_values:
            SupplierInfo.check_access("create")
            SupplierInfo.create(create_values)

    def action_apply(self):
        self.ensure_one()
        self._check_order()
        if not self.preview_valid:
            raise UserError(_("Validate the import successfully before applying it."))
        if self.validation_token != self._configuration_digest():
            raise UserError(_("The file or import configuration changed; generate a new preview."))
        if self.order_snapshot != self._order_line_snapshot():
            raise UserError(_("The purchase order lines changed; generate a new preview."))
        self._validate_mapping()
        parsed = self._parse_file()
        normalized_rows = self._validate_rows(parsed)
        errors = [
            _("Row %(row)s: %(message)s", row=row["source_row"], message="; ".join(row["errors"]))
            for row in normalized_rows
            if row["errors"]
        ]
        if errors:
            raise UserError("\n".join(errors))
        self._check_order()
        PurchaseLine = self.env["purchase.order.line"]
        PurchaseLine.check_access("create")
        product_cache = {}
        line_commands = []
        normalized_products = []
        for normalized in normalized_rows:
            product = normalized["product"] or product_cache.get(normalized["plan_key"])
            if not product:
                product = (
                    self._create_dynamic_variant(normalized)
                    if normalized["plan_kind"] == "variant"
                    else self._create_standalone_product(normalized)
                )
                product_cache[normalized["plan_key"]] = product
            normalized_products.append((normalized, product))
            line_values = {
                "product_id": product.id,
                "product_qty": normalized["quantity"],
                "product_uom_id": normalized["uom"].id,
                "price_unit": normalized["price"],
            }
            if normalized.get("description"):
                line_values["name"] = normalized["description"]
            line_commands.append(Command.create(line_values))
        self._upsert_supplierinfos(normalized_products)
        self.order_id.with_company(self.order_id.company_id).write({
            "order_line": line_commands,
        })
        return {"type": "ir.actions.act_window_close"}

    def _invalidate_preview(self):
        self.with_context(purchase_import_keep_preview=True).write({
            "preview_valid": False,
            "validation_token": False,
            "state": "mapping",
        })

    def _action_reopen(self):
        action = self.env["ir.actions.actions"]._for_xml_id(
            "purchase_import.purchase_order_line_import_wizard_action"
        )
        action.update({"res_id": self.id, "target": "new"})
        return action


class PurchaseOrderLineImportMapping(models.TransientModel):
    _name = "purchase.order.line.import.mapping"
    _description = "Purchase Order Line Import Mapping"
    _order = "sequence, source_index, id"

    wizard_id = fields.Many2one(
        "purchase.order.line.import.wizard", required=True, ondelete="cascade", index=True
    )
    sequence = fields.Integer(default=10)
    source_index = fields.Integer(required=True)
    source_column = fields.Char(required=True)
    target = fields.Selection(TARGET_SELECTION)
    attribute_id = fields.Many2one("product.attribute", ondelete="restrict")

    @api.constrains("target", "attribute_id")
    def _check_attribute_target(self):
        for mapping in self:
            if mapping.target == "attribute" and not mapping.attribute_id:
                raise ValidationError(_("Select an attribute for an attribute mapping."))
            if mapping.target != "attribute" and mapping.attribute_id:
                raise ValidationError(_("Clear the attribute on a non-attribute mapping."))

    @api.model_create_multi
    def create(self, vals_list):
        mappings = super().create(vals_list)
        if not self.env.context.get("purchase_import_keep_preview"):
            mappings.wizard_id._invalidate_preview()
        return mappings

    def write(self, vals):
        result = super().write(vals)
        if not self.env.context.get("purchase_import_keep_preview"):
            self.wizard_id._invalidate_preview()
        return result

    def unlink(self):
        wizards = self.wizard_id
        result = super().unlink()
        if not self.env.context.get("purchase_import_keep_preview"):
            wizards._invalidate_preview()
        return result


class PurchaseOrderLineImportPreview(models.TransientModel):
    _name = "purchase.order.line.import.preview"
    _description = "Purchase Order Line Import Preview"
    _order = "source_row, id"

    wizard_id = fields.Many2one(
        "purchase.order.line.import.wizard", required=True, ondelete="cascade", index=True
    )
    source_row = fields.Integer(readonly=True)
    status = fields.Selection(
        [("resolved", "Resolved"), ("create", "Create"), ("error", "Error")],
        readonly=True,
    )
    product_id = fields.Many2one("product.product", readonly=True)
    planned_action = fields.Char(readonly=True)
    product_name = fields.Char(readonly=True)
    quantity = fields.Float(readonly=True)
    uom_id = fields.Many2one("uom.uom", readonly=True)
    price = fields.Float(readonly=True)
    min_qty = fields.Float(readonly=True)
    supplier_code = fields.Char(readonly=True)
    default_code = fields.Char(readonly=True)
    barcode = fields.Char(readonly=True)
    attribute_values = fields.Char(readonly=True)
    error_message = fields.Text(readonly=True)
